# research/dislocation_study/03_dislocation_study.py
"""
STEP 3 -- Dislocation study + Excel report (multi-sheet).

Question: does the Stage 6/7 breakout have an edge ONLY when price is stretched a
certain % BELOW the 200 EMA after a Stage 2 (sharp downtrend)? And what distance
(or distance band) is best across periods?

Trades come from step 2 (exact backtest-engine replica, one position per ticker).
Everything is NET of commission (study_config.COMMISSION_PER_SHARE_RT) and R is
winsorized at +WINSOR_R so a single lottery ticket cannot drive a result.

Periods: TRAIN 2022-01..2023-12 | TEST 2024-01..2026-01 | OOS26 2026-02.. (new data).
A candidate is 'robust' only if mean net R > 0 in TRAIN and TEST (and OOS26 when it
has enough trades) with minimum trade counts. Rankings use min(train, test) so one
good period cannot carry a candidate.  Bootstrap CIs resample whole signal DATES
(clusters) because same-day signals are correlated.

Output: output/Dislocation_Study_Report.xlsx  (sheets listed in the 'Index' sheet)

Run:  python 03_dislocation_study.py
"""
from __future__ import annotations

import warnings

import numpy as np
import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import study_config as C

warnings.filterwarnings("ignore")
PERIODS = ["TRAIN", "TEST", "OOS26"]


# ── Data ─────────────────────────────────────────────────────────────────────
def load_trades() -> pd.DataFrame:
    T = pd.read_parquet(C.SIM_TRADES)
    T = T[T.taken & ~T.exit_reason.isin(["end_of_data", "stage9_exit_eod", "time_stop_eod"])].copy()
    T["net_r"] = T.pnl_r - C.COMMISSION_PER_SHARE_RT / T.stop_dist - C.EXTRA_SLIPPAGE_R
    T["w"] = T.net_r.clip(upper=C.WINSOR_R)
    sd = T.signal_date
    T["period"] = np.select(
        [sd <= pd.Timestamp(C.LEGACY_END),
         (sd >= pd.Timestamp(C.TRAIN[0])) & (sd <= pd.Timestamp(C.TRAIN[1])),
         (sd >= pd.Timestamp(C.TEST[0])) & (sd <= pd.Timestamp(C.TEST[1])),
         sd >= pd.Timestamp(C.OOS26[0])],
        ["LEGACY", "TRAIN", "TEST", "OOS26"], default="GAP")
    T["year"] = sd.dt.year
    T["stack_none"] = True
    T["stack_e10e20"] = T.e10 > T.e20
    T["stack_full"] = (T.e10 > T.e20) & (T.e20 > T.e50)
    return T.reset_index(drop=True)


def metrics(w: np.ndarray, net: np.ndarray, tick: np.ndarray, dates: np.ndarray) -> dict:
    n = len(w)
    if n == 0:
        return {"n": 0, "tickers": 0, "days": 0, "win_pct": np.nan, "mean_net_R": np.nan,
                "median_net_R": np.nan, "pf": np.nan}
    pos, neg = w[w > 0].sum(), -w[w < 0].sum()
    return {"n": n, "tickers": len(set(tick)), "days": len(set(dates)),
            "win_pct": float((net > 0).mean() * 100), "mean_net_R": float(w.mean()),
            "median_net_R": float(np.median(net)), "pf": float(pos / neg) if neg > 0 else np.nan}


def mask_of(T, x=None, lo=None, hi=None, n=None, stack="none") -> np.ndarray:
    m = np.ones(len(T), bool)
    d = T.dist200.values
    if x is not None:
        m &= d <= -x
    if lo is not None:
        m &= (d > lo) & (d <= hi)
    if n is not None:
        m &= T.days_since_s2.values <= n
    col = {"none": "stack_none", "e10>e20": "stack_e10e20", "e10>e20>e50": "stack_full"}[stack]
    m &= T[col].values
    return m & ~np.isnan(d)


def by_period(T, m: np.ndarray, prefix: str = "") -> dict:
    out = {}
    for p in PERIODS + ["ALL22+"]:
        pm = m & ((T.period.values == p) if p != "ALL22+" else np.isin(T.period.values, PERIODS))
        r = metrics(T.w.values[pm], T.net_r.values[pm], T.ticker.values[pm], T.signal_date.values[pm])
        for k in ("n", "win_pct", "mean_net_R", "pf"):
            out[f"{p}_{k}"] = r[k]
        if p == "ALL22+":
            out["tickers"] = r["tickers"]
    return out


def robust_flags(r: dict) -> dict:
    ok = (r["TRAIN_n"] >= C.MIN_N and r["TEST_n"] >= C.MIN_N
          and r["TRAIN_mean_net_R"] > 0 and r["TEST_mean_net_R"] > 0)
    oos_ok = (r["OOS26_n"] < C.MIN_N_OOS) or (r["OOS26_mean_net_R"] > 0)
    mm = np.nanmin([r["TRAIN_mean_net_R"], r["TEST_mean_net_R"]]) if ok or r["TRAIN_n"] and r["TEST_n"] else np.nan
    return {"min_train_test_R": mm, "robust": bool(ok and oos_ok),
            "oos_checked": bool(r["OOS26_n"] >= C.MIN_N_OOS)}


# ── Studies ──────────────────────────────────────────────────────────────────
def study_bands(T) -> pd.DataFrame:
    edges = [-1, -.5, -.4, -.3, -.25, -.2, -.15, -.1, -.05, 0, .05, .1, .2, .4, 100]
    rows = []
    for lo, hi in zip(edges[:-1], edges[1:]):
        m = mask_of(T, lo=lo, hi=hi)
        rows.append({"dist200_band": f"({lo:+.0%}, {hi:+.0%}]" if hi < 100 else f"> {lo:+.0%}", **by_period(T, m)})
    return pd.DataFrame(rows)


def study_matrix(T):
    de = [-1, -.3, -.2, -.1, 0, .1, 100]
    dl = ["< -30%", "-30..-20%", "-20..-10%", "-10..0%", "0..+10%", "> +10%"]
    se = [-1, 20, 60, 120, 260, 1e9]
    sl = ["<=20", "21-60", "61-120", "121-260", ">260"]
    M = pd.DataFrame(index=dl, columns=sl, dtype=float)
    N = pd.DataFrame(index=dl, columns=sl, dtype=float)
    ok = T.period.isin(PERIODS) & T.days_since_s2.notna()
    for i, dlab in enumerate(dl):
        for j, slab in enumerate(sl):
            m = (ok & (T.dist200 > de[i]) & (T.dist200 <= de[i + 1])
                 & (T.days_since_s2 > se[j]) & (T.days_since_s2 <= se[j + 1])).values
            M.iloc[i, j] = T.w.values[m].mean() if m.sum() >= 30 else np.nan
            N.iloc[i, j] = m.sum()
    M.index.name, N.index.name = "dist200 \\ days since Stage 2", "dist200 \\ days since Stage 2"
    return M.reset_index(), N.reset_index()


def study_threshold_grid(T) -> pd.DataFrame:
    rows = []
    for stack in C.STACKS:
        for x in [None] + C.X_BELOW:
            for n in C.N_S2:
                r = by_period(T, mask_of(T, x=x, n=n, stack=stack))
                rows.append({"stack": stack, "min_below_ema200": x, "stage2_within_bars": n, **r, **robust_flags(r)})
    G = pd.DataFrame(rows)
    # plateau score: mean of min_train_test_R over the 8 neighbouring cells (same stack)
    xs, ns = [None] + C.X_BELOW, C.N_S2
    lookup = {(r.stack, r.min_below_ema200, r.stage2_within_bars): r.min_train_test_R for r in G.itertuples()}
    plat = []
    for r in G.itertuples():
        i = xs.index(r.min_below_ema200 if r.min_below_ema200 == r.min_below_ema200 else None)
        j = ns.index(r.stage2_within_bars if r.stage2_within_bars == r.stage2_within_bars else None)
        vals = []
        for di in (-1, 0, 1):
            for dj in (-1, 0, 1):
                if (di, dj) == (0, 0) or not (0 <= i + di < len(xs) and 0 <= j + dj < len(ns)):
                    continue
                v = lookup.get((r.stack, xs[i + di], ns[j + dj]))
                if v is not None and v == v:
                    vals.append(v)
        plat.append(np.mean(vals) if vals else np.nan)
    G["neighbour_avg_min_R"] = plat
    return G


def study_band_grid(T) -> pd.DataFrame:
    rows = []
    for stack in C.STACKS:
        for lo in C.BAND_LO:
            for hi in C.BAND_HI:
                if hi <= lo + 0.049:
                    continue
                for n in C.BAND_N:
                    r = by_period(T, mask_of(T, lo=lo, hi=hi, n=n, stack=stack))
                    rows.append({"stack": stack, "band_low": lo, "band_high": hi,
                                 "stage2_within_bars": n, **r, **robust_flags(r)})
    return pd.DataFrame(rows)


def cluster_bootstrap(T, m, period_mask) -> tuple:
    pm = m & period_mask
    if pm.sum() < 20:
        return (np.nan, np.nan)
    df = pd.DataFrame({"d": T.signal_date.values[pm], "w": T.w.values[pm]})
    g = df.groupby("d").w.agg(["sum", "count"]).values
    rng = np.random.default_rng(C.SEED)
    idx = rng.integers(0, len(g), size=(C.BOOTSTRAP_REPS, len(g)))
    means = g[idx, 0].sum(1) / g[idx, 1].sum(1)
    return (float(np.percentile(means, 2.5)), float(np.percentile(means, 97.5)))


def top_candidates(T, B, Bd, k=25):
    b = B[B.robust].assign(source="B_threshold", spec=lambda d: d.apply(
        lambda r: f"dist200<=-{r.min_below_ema200:.0%}" if r.min_below_ema200 == r.min_below_ema200 else "no distance filter", axis=1))
    d = Bd[Bd.robust].assign(source="C_band", spec=lambda d: d.apply(
        lambda r: f"dist200 in ({r.band_low:+.0%},{r.band_high:+.0%}]", axis=1))
    cand = pd.concat([b, d], ignore_index=True)
    if cand.empty:
        return pd.DataFrame(), []
    cand = cand.sort_values("min_train_test_R", ascending=False).head(k).reset_index(drop=True)
    pmask = {p: (T.period.values == p) for p in PERIODS}
    pmask["ALL22+"] = np.isin(T.period.values, PERIODS)
    masks, rows = [], []
    for r in cand.itertuples():
        if r.source == "B_threshold":
            x = None if r.min_below_ema200 != r.min_below_ema200 else r.min_below_ema200
            n = None if r.stage2_within_bars != r.stage2_within_bars else r.stage2_within_bars
            m = mask_of(T, x=x, n=n, stack=r.stack)
        else:
            n = None if r.stage2_within_bars != r.stage2_within_bars else r.stage2_within_bars
            m = mask_of(T, lo=r.band_low, hi=r.band_high, n=n, stack=r.stack)
        masks.append(m)
        ci_test = cluster_bootstrap(T, m, pmask["TEST"])
        ci_oos = cluster_bootstrap(T, m, pmask["OOS26"])
        ci_all = cluster_bootstrap(T, m, pmask["ALL22+"])
        rows.append({"rank": len(rows) + 1, "source": r.source, "spec": r.spec, "stack": r.stack,
                     "stage2_within_bars": n, "tickers": r.tickers,
                     "TRAIN_n": r.TRAIN_n, "TRAIN_mean_net_R": r.TRAIN_mean_net_R,
                     "TEST_n": r.TEST_n, "TEST_mean_net_R": r.TEST_mean_net_R,
                     "OOS26_n": r.OOS26_n, "OOS26_mean_net_R": r.OOS26_mean_net_R,
                     "ALL_n": r._asdict()["_" + str(list(cand.columns).index("ALL22+_n"))] if False else cand.loc[r.Index, "ALL22+_n"],
                     "ALL_mean_net_R": cand.loc[r.Index, "ALL22+_mean_net_R"],
                     "ALL_pf": cand.loc[r.Index, "ALL22+_pf"],
                     "min_train_test_R": r.min_train_test_R,
                     "TEST_CI95_low": ci_test[0], "TEST_CI95_high": ci_test[1],
                     "OOS26_CI95_low": ci_oos[0], "OOS26_CI95_high": ci_oos[1],
                     "ALL_CI95_low": ci_all[0], "ALL_CI95_high": ci_all[1],
                     "ALL_CI_excludes_0": bool(ci_all[0] > 0) if ci_all[0] == ci_all[0] else False})
    return pd.DataFrame(rows), masks


def yearly_and_concentration(T, top, masks, k=5):
    ys, cs = [], []
    for i in range(min(k, len(masks))):
        m = masks[i]
        spec = f"#{i + 1} {top.loc[i, 'spec']} | {top.loc[i, 'stack']} | S2<= {top.loc[i, 'stage2_within_bars']}"
        for y in sorted(T.year[m].unique()):
            ym = m & (T.year.values == y)
            r = metrics(T.w.values[ym], T.net_r.values[ym], T.ticker.values[ym], T.signal_date.values[ym])
            ys.append({"candidate": spec, "year": int(y), **{k2: r[k2] for k2 in ("n", "win_pct", "mean_net_R", "pf")}})
        sub = T[m & np.isin(T.period.values, PERIODS)]
        if len(sub):
            pos = sub[sub.w > 0]
            by_t = pos.groupby("ticker").w.sum().sort_values(ascending=False)
            cs.append({"candidate": spec, "trades": len(sub), "tickers": sub.ticker.nunique(),
                       "top10_ticker_share_of_positive_R": float(by_t.head(10).sum() / by_t.sum()) if len(by_t) else np.nan,
                       "largest_single_trade_share": float(pos.w.max() / pos.w.sum()) if len(pos) else np.nan,
                       "top3_tickers": ", ".join(by_t.head(3).index)})
    return pd.DataFrame(ys), pd.DataFrame(cs)


def study_coverage() -> pd.DataFrame:
    """Per-month continuity proof: signals, taken trades, closed vs still-open (no gap 2022 -> now)."""
    R = pd.read_parquet(C.SIM_TRADES, columns=["ticker", "signal_date", "taken", "exit_reason"])
    R = R[R.signal_date >= pd.Timestamp(C.TRAIN[0])]
    R["month"] = R.signal_date.dt.to_period("M").astype(str)
    g = R.groupby("month")
    return pd.DataFrame({
        "signals": g.size(),
        "taken": g.taken.sum(),
        "tickers_signalling": g.ticker.nunique(),
        "closed": g.apply(lambda d: int((d.taken & ~d.exit_reason.isin(["end_of_data", "stage9_exit_eod", "time_stop_eod"])).sum())),
        "still_open": g.apply(lambda d: int((d.taken & d.exit_reason.isin(["end_of_data", "stage9_exit_eod", "time_stop_eod"])).sum())),
    }).reset_index()


def study_baselines(T) -> pd.DataFrame:
    rows = []
    def add(label, m):
        rows.append({"slice": label, **by_period(T, m)})
    add("All signals (taken)", np.ones(len(T), bool))
    add("Stage 6 breakout", (T.stage.values == 6))
    add("Stage 7 confirmed", (T.stage.values == 7))
    for ps in sorted(T.prev_stage.dropna().unique()):
        add(f"prev stage = {int(ps)}", T.prev_stage.values == ps)
    add("EMA stack e10>e20", T.stack_e10e20.values)
    add("EMA stack e10>e20>e50", T.stack_full.values)
    add("dist200 < 0 (below EMA200)", (T.dist200.values < 0))
    add("dist200 >= 0 (above EMA200)", (T.dist200.values >= 0))
    return pd.DataFrame(rows)


# ── Excel formatting ─────────────────────────────────────────────────────────
HDR = PatternFill("solid", start_color="1F3864")
HFONT = Font(bold=True, color="FFFFFF")


def pretty(ws, df, pct_cols=(), r_cols=(), int_cols=(), widths=None, scale_cols=()):
    ws.freeze_panes = "A2"
    for c in range(1, df.shape[1] + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HDR, HFONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(df.columns[c - 1], max(12, min(34, len(str(df.columns[c - 1])) + 2)))
    ws.row_dimensions[1].height = 42
    for ci, name in enumerate(df.columns, 1):
        fmt = None
        if any(name.endswith(s) for s in ("win_pct",)):
            fmt = '0.0'
        if name in pct_cols or any(s in name for s in ("share", "min_below", "band_low", "band_high")):
            fmt = "0.0%"
        if "mean_net_R" in name or "median" in name or name in r_cols or "min_train_test" in name or "neighbour" in name or "CI95" in name:
            fmt = "0.000"
        if name.endswith("_pf") or name == "pf":
            fmt = "0.00"
        if name.endswith("_n") or name in int_cols or name in ("tickers", "days", "trades", "year"):
            fmt = "#,##0"
        if fmt:
            for r in range(2, df.shape[0] + 2):
                ws.cell(row=r, column=ci).number_format = fmt
    for name in scale_cols:
        if name in df.columns and df.shape[0]:
            col = get_column_letter(list(df.columns).index(name) + 1)
            ws.conditional_formatting.add(f"{col}2:{col}{df.shape[0] + 1}", ColorScaleRule(
                start_type="num", start_value=-0.3, start_color="F8696B",
                mid_type="num", mid_value=0, mid_color="FFFFFF",
                end_type="num", end_value=0.3, end_color="63BE7B"))


def main() -> None:
    T = load_trades()
    print(f"[INFO] closed taken trades: {len(T):,}; periods: {T.period.value_counts().to_dict()}")
    A = study_bands(T)
    M, N = study_matrix(T)
    B = study_threshold_grid(T)
    Bd = study_band_grid(T)
    top, masks = top_candidates(T, B, Bd)
    Y, CON = (yaml_empty(), yaml_empty()) if top.empty else yearly_and_concentration(T, top, masks)
    BASE = study_baselines(T)
    COV = study_coverage()

    parity = None
    pp = C.OUT / "00_parity_validation.csv"
    if pp.exists():
        parity = pd.read_csv(pp)

    n_tests = len(B) + len(Bd)
    n_robust = int(B.robust.sum() + Bd.robust.sum())
    idx = pd.DataFrame({"sheet": [
        "Summary", "Baselines", "A_Distance_Bands", "A2_Matrix_MeanR", "A2_Matrix_N",
        "B_Threshold_Grid", "C_Band_Grid", "D_Top_Candidates", "E_Yearly", "F_Concentration",
        "Stages", "Coverage", "Parity"],
        "what_it_shows": [
            "Headline findings and caveats", "Reference slices (stage, prev stage, EMA stack, above/below EMA200)",
            "Net R by signed distance of close from EMA200 (negative = below), per period",
            "Mean net R: distance below EMA200 x days since Stage 2", "Trade counts for the matrix",
            "Grid: price at least X below EMA200 x Stage 2 within N bars x EMA stack",
            "Grid: distance BAND (lo,hi] x Stage 2 recency x EMA stack",
            "Robust candidates ranked by min(train,test) with date-cluster bootstrap 95% CI",
            "Year-by-year for the top 5 candidates", "Ticker concentration for the top 5",
            "Stage definitions (from production code)",
            "Monthly signals/trades 2022 -> now (continuity proof, no gap between backtest and Work PC data)",
            "Proof that the simulator reproduces the original backtest"]})

    summ = pd.DataFrame({"item": [
        "Closed taken trades analysed", "Periods", "Net of commission ($/share RT)", "R winsorized at",
        "Grid combinations tested (B + C)", "Combinations flagged robust", "Expected false positives",
        "Best candidate", "Interpretation rule"],
        "value": [
            f"{len(T):,}", f"TRAIN {C.TRAIN[0]}..{C.TRAIN[1]} | TEST {C.TEST[0]}..{C.TEST[1]} | OOS26 {C.OOS26[0]}..",
            C.COMMISSION_PER_SHARE_RT, f"+{C.WINSOR_R}R", n_tests, n_robust,
            "With this many tests some positives appear by chance alone - trust plateaus and CI that excludes 0, not single cells",
            (f"{top.loc[0, 'spec']} | stack {top.loc[0, 'stack']} | S2<= {top.loc[0, 'stage2_within_bars']} | "
             f"min(train,test)={top.loc[0, 'min_train_test_R']:.3f}R, ALL CI95 [{top.loc[0, 'ALL_CI95_low']:.3f}, {top.loc[0, 'ALL_CI95_high']:.3f}]") if not top.empty else "none passed the robust test",
            "A real edge should (1) be positive in TRAIN, TEST and OOS26, (2) sit on a plateau of neighbouring cells, (3) have a bootstrap CI above 0, (4) not depend on a few tickers."]})

    stages = pd.DataFrame({
        "stage": [1, 2, 3, 4, 5, 6, 7, 8, 9],
        "name": ["Not Eligible", "Sharp Downtrend", "Downtrend", "Below Zone", "Lower Zone",
                 "Breakout", "Breakout Confirmed", "In-Zone", "In-Zone (Fading)"],
        "rule (checked in this order: 2,3,4,5,7,6,9,8)": [
            "fallback, or < 260 bars of history",
            "below EMA200, EMA10<20<50<200, close < prior 20d low, volume expansion",
            "below EMA200 and bearish EMA stack",
            "below EMA200 and close <= lower Bollinger band",
            "below EMA200 and close between lower and mid Bollinger band",
            "close > prior 20d high and EMA10 > EMA20",
            "close > prior 20d high, close > EMA50, EMA10 > EMA20",
            "above EMA200",
            "above EMA200 and close < EMA10 (also the exit trigger)"]})

    C.OUT.mkdir(parents=True, exist_ok=True)
    path = C.OUT / "Dislocation_Study_Report.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        sheets = [("Index", idx), ("Summary", summ), ("Baselines", BASE), ("A_Distance_Bands", A),
                  ("A2_Matrix_MeanR", M), ("A2_Matrix_N", N), ("B_Threshold_Grid", B),
                  ("C_Band_Grid", Bd), ("D_Top_Candidates", top), ("E_Yearly", Y),
                  ("F_Concentration", CON), ("Stages", stages), ("Coverage", COV)]
        if parity is not None:
            sheets.append(("Parity", parity))
        for name, df in sheets:
            df.to_excel(xw, sheet_name=name, index=False)
        wb = xw.book
        for name, df in sheets:
            ws = wb[name]
            pretty(ws, df, scale_cols=[c for c in df.columns if "mean_net_R" in c or c == "min_train_test_R"],
                   widths={"item": 36, "value": 110, "what_it_shows": 90, "sheet": 24, "spec": 38,
                           "candidate": 70, "slice": 34, "dist200_band": 22, "rule (checked in this order: 2,3,4,5,7,6,9,8)": 80})
        if "A2_Matrix_MeanR" in wb.sheetnames:
            ws = wb["A2_Matrix_MeanR"]
            for row in ws.iter_rows(min_row=2, min_col=2):
                for cell in row:
                    cell.number_format = "0.000"
            ws.column_dimensions["A"].width = 30
    print(f"[OK] report -> {path}")
    if not top.empty:
        print(top[["rank", "spec", "stack", "stage2_within_bars", "TRAIN_mean_net_R", "TEST_mean_net_R",
                   "OOS26_mean_net_R", "ALL_n", "ALL_CI95_low", "ALL_CI95_high"]].head(10).round(3).to_string(index=False))
    else:
        print("[RESULT] no combination passed the robust test (positive net R in TRAIN and TEST with min trade counts).")


def yaml_empty() -> pd.DataFrame:
    return pd.DataFrame()


if __name__ == "__main__":
    main()
