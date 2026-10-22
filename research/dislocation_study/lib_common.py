# research/dislocation_study/lib_common.py
"""Shared helpers for steps 04 (cost stress) and 05 (portfolio sim). Read-only on all data."""
from __future__ import annotations

import glob

import numpy as np
import pandas as pd
import pyarrow.parquet as pq

import study_config as C

PERIODS = ["TRAIN", "TEST", "OOS26"]


def period_of(sd: pd.Series) -> np.ndarray:
    return np.select(
        [sd <= pd.Timestamp(C.LEGACY_END),
         (sd >= pd.Timestamp(C.TRAIN[0])) & (sd <= pd.Timestamp(C.TRAIN[1])),
         (sd >= pd.Timestamp(C.TEST[0])) & (sd <= pd.Timestamp(C.TEST[1])),
         sd >= pd.Timestamp(C.OOS26[0])],
        ["LEGACY", "TRAIN", "TEST", "OOS26"], default="GAP")


def add_adv20(T: pd.DataFrame) -> pd.DataFrame:
    """20-day average dollar volume at the signal bar (point-in-time), read shard by shard."""
    need = T[["ticker", "signal_date"]].drop_duplicates()
    parts = []
    for f in sorted(glob.glob(str(C.SHARDS / "s*.parquet"))):
        s = pq.read_table(f, columns=["ticker", "date", "close", "volume"]).to_pandas()
        s = s.sort_values(["ticker", "date"])
        s["dv"] = s.close * s.volume
        s["adv20"] = s.groupby("ticker").dv.transform(lambda x: x.rolling(20, min_periods=10).mean())
        s = s[["ticker", "date", "adv20"]].rename(columns={"date": "signal_date"})
        parts.append(s.merge(need, on=["ticker", "signal_date"]))
    A = pd.concat(parts, ignore_index=True)
    return T.merge(A, on=["ticker", "signal_date"], how="left")


def load_all_signals(closed_only: bool = False) -> pd.DataFrame:
    """ALL simulated signals (taken or not) from PORT_START on, with costs-ready columns.
    Rows flagged `open_at_end` were still open at the last data bar (valued at last close)."""
    T = pd.read_parquet(C.SIM_TRADES)
    T = T[T.entry_date >= pd.Timestamp(C.PORT_START)].copy()
    T["open_at_end"] = T.exit_reason.isin(["end_of_data", "stage9_exit_eod", "time_stop_eod"])
    if closed_only:
        T = T[~T.open_at_end]
    T["period"] = period_of(T.signal_date)
    T["year"] = T.signal_date.dt.year
    T = add_adv20(T)
    return T.reset_index(drop=True)


def spec_mask(T: pd.DataFrame, spec: tuple) -> np.ndarray:
    x, n = spec
    m = ~np.isnan(T.dist200.values)
    if x is not None:
        m &= T.dist200.values <= -x
    if n is not None:
        m &= T.days_since_s2.values <= n
    return m


def net_r(T: pd.DataFrame, slip: float) -> np.ndarray:
    """Net R per trade: commission + swap + extra slippage (all in R)."""
    return (T.pnl_r - C.COMMISSION_PER_SHARE_RT / T.stop_dist - C.SWAP_R_PER_TRADE - slip).values


def stats(net: np.ndarray) -> dict:
    n = len(net)
    if n == 0:
        return {"n": 0, "win_pct": np.nan, "mean_R": np.nan, "pf": np.nan}
    w = np.minimum(net, C.WINSOR_R)
    pos, neg = w[w > 0].sum(), -w[w < 0].sum()
    return {"n": n, "win_pct": float((net > 0).mean() * 100), "mean_R": float(w.mean()),
            "pf": float(pos / neg) if neg > 0 else np.nan}


def cluster_ci(net: np.ndarray, dates: np.ndarray, reps: int = C.BOOTSTRAP_REPS, seed: int = C.SEED):
    """95% CI of mean winsorized R, resampling signal-dates (clusters)."""
    if len(net) < 30:
        return np.nan, np.nan
    w = np.minimum(net, C.WINSOR_R)
    codes, uniq = pd.factorize(dates)
    k = len(uniq)
    s = np.bincount(codes, weights=w, minlength=k)
    c = np.bincount(codes, minlength=k).astype(float)
    rng = np.random.default_rng(seed)
    idx = rng.integers(0, k, size=(reps, k))
    means = s[idx].sum(1) / c[idx].sum(1)
    return float(np.percentile(means, 2.5)), float(np.percentile(means, 97.5))


def style_workbook(path, widths: dict | None = None) -> None:
    """Uniform look: bold dark header, freeze panes, auto widths, number formats."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "B2"
        for i, col in enumerate(ws.columns, 1):
            mx = max((len(str(c.value)) if c.value is not None else 0) for c in list(col)[:200])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, mx + 2), 48)
            for c in list(col)[1:]:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.000" if abs(c.value) < 50 else "#,##0"
    wb.save(path)
